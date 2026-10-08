# -*- coding: utf-8 -*-
"""Scan every generated PPTX for HTML entities the slide renderer cannot decode.
Slides take raw Unicode; notes take entities. A leak shows up as literal text."""
import glob, os, re, sys
from pptx import Presentation

# Entities the deck engine DOES decode (uc.py _ENTITIES) are not leaks.
OK = {"&amp;", "&lt;", "&gt;", "&nbsp;", "&mdash;", "&ndash;", "&middot;",
      "&minus;", "&rarr;", "&eacute;", "&quot;"}
ENT = re.compile(r"&(?:#\d{2,5}|[a-zA-Z]+);")

bad = 0
for f in sorted(glob.glob("Courses/*/Slides/*.pptx")):
    for i, s in enumerate(Presentation(f).slides, 1):
        texts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                texts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    texts += [c.text for c in r.cells]
        for t in texts:
            for m in set(ENT.findall(t)) - OK:
                bad += 1
                print("%s  slide %d  %s" % (f, i, m))
print("%d leaked entities" % bad)
sys.exit(1 if bad else 0)
