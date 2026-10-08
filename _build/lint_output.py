# -*- coding: utf-8 -*-
"""Scan every generated PPTX *and* PDF for undecoded HTML entities.

`lint_slides.py` only ever looked at the decks, which is how a batch of
leaks survived in the notes: `Doc.callout()` uppercased its label before
rendering, and entity names are case-sensitive, so "&times;" became
"&TIMES;" and no parser would decode it. Anything that renders text is
worth scanning, so this scans both.

Exits non-zero on any leak, so it can gate a build.
"""
import glob, os, re, sys

# Globs resolve against the repository root, not the working directory.
# These used to be bare relative paths, so running this from _build
# scanned nothing and still printed a pass.
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def _real(paths):
    """Drop Office lock files (~$foo.pptx); they are not documents."""
    return [p for p in paths
            if not os.path.basename(p).startswith("~$")]
import pymupdf
from pptx import Presentation

# Entities the deck engine decodes itself (uc.py _NAMED) are not leaks in
# slides; reportlab decodes the standard set in the PDF path.
DECK_OK = {"&amp;", "&lt;", "&gt;", "&nbsp;", "&mdash;", "&ndash;",
           "&middot;", "&minus;", "&rarr;", "&eacute;", "&quot;"}
ENT = re.compile(r"&(?:#\d{2,5}|[a-zA-Z][a-zA-Z0-9]{1,9});")

bad = []


def deck_text(path):
    for i, s in enumerate(Presentation(path).slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame:
                yield i, sh.text_frame.text
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        yield i, c.text


def main():
    decks = _real(sorted(glob.glob(os.path.join(REPO, "Courses/*/Slides/*.pptx"))))
    for f in decks:
        for i, t in deck_text(f):
            for m in set(ENT.findall(t)) - DECK_OK:
                bad.append("%s  slide %d  %s" % (f, i, m))

    pdfs = (sorted(glob.glob(os.path.join(REPO, "Courses/*/Notes/*.pdf")))
            + sorted(glob.glob(os.path.join(REPO, "Courses/*/*.pdf")))
            + sorted(glob.glob(os.path.join(REPO, "00_Program/*.pdf"))))
    for f in pdfs:
        d = pymupdf.open(f)
        for i in range(d.page_count):
            for m in set(ENT.findall(d[i].get_text())):
                bad.append("%s  page %d  %s" % (f, i + 1, m))
        d.close()

    for b in bad[:60]:
        print(b)
    if len(bad) > 60:
        print("... and %d more" % (len(bad) - 60))
    print("%d decks + %d PDFs scanned, %d leaked entities"
          % (len(decks), len(pdfs), len(bad)))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
